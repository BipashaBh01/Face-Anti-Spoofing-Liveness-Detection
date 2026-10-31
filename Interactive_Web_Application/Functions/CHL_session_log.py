# ============================================================================
#  CHL_session_log.py — append one row per completed session to the study log
# ============================================================================
#  One completed session produces one row in an Excel (.xlsx) workbook. Each
#  attempt within that session is recorded inside the challenge columns,
#  labelled by attempt number and result. The workbook is created with a
#  header the first time and appended to thereafter.
# ============================================================================

from __future__ import annotations

# datetime stamps each row with the date and time of writing.
from datetime import datetime

# Path handles the file location and creates the parent folder when missing.
from pathlib import Path

# openpyxl reads and writes native Excel .xlsx workbooks.
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font

# Location of the study log, taken from the central config.
from Functions.CHL_config import SESSION_LOG_PATH as XLSX_PATH


# The fixed column order. The header row and every written row follow this order.
COLUMNS = [
    "session_id",          # unique id of the session (date_time_counter_participant)
    "participant_code",    # participant code entered on the setup screen
    "age_group",           # participant age band
    "gender",              # participant gender
    "ethnicity",           # participant ethnicity
    "visual_aids",         # whether visual aids (glasses/contacts) were worn
    "datetime",            # date and time the row was written
    "condition",           # the test condition for this session
    "challenge_actions",   # the three actions issued, per attempt
    "per_action_outcome",  # per-action detected/not-detected result, per attempt
    "per_action_detail",   # per-action measured detail (score/angle/frames), per attempt
    "challenge_verdict",   # overall challenge result: PASSED or FAILED
    "passive_verdict",     # passive model verdict: LIVE or SPOOF
    "passive_live_prob",   # passive model live probability
    "passive_spoof_prob",  # passive model spoof probability
    "final_result",        # combined outcome across the passive and active layers
    "final_user_message",  # the message shown on screen to the user
    "device_browser",      # browser/device user-agent string
]


def _fmt_attempts(attempts):
    """Format the list of attempt records into three multi-line strings.

    Each attempt is a dict with:
      actions  : list[str]  — the three action keys issued
      outcomes : list[str]  — per-action detected / not detected
      details  : list[str]  — per-action measured detail
      result   : str        — 'passed' or 'failed'
      swapped  : bool        — True when the attempt was an identity swap
    """

    # With no attempts (for example a session that never reached the challenge),
    # return three empty cells.
    if not attempts:
        return "", "", ""

    # Collect the formatted fragment for each attempt, one per list.
    a_parts, o_parts, d_parts = [], [], []

    # Number the attempts from one and format each in turn.
    for i, att in enumerate(attempts, start=1):

        # Label carrying the attempt number and its result, e.g. "attempt2 (failed)".
        tag = f"attempt{i} ({att.get('result', 'failed')})"

        # The actions issued for this attempt, comma-separated.
        a_parts.append(f"{tag}: " + ", ".join(att.get("actions", [])))

        # An identity-swapped attempt is recorded plainly as such in both the
        # outcome and detail columns, since no normal action result applies.
        if att.get("swapped"):
            o_parts.append(f"{tag}: identity swapped")
            d_parts.append(f"{tag}: identity swapped")

        # A normal attempt records its per-action outcomes and measured details.
        else:
            o_parts.append(f"{tag}: " + ", ".join(att.get("outcomes", [])))
            d_parts.append(f"{tag}: " + "; ".join(att.get("details", [])))

    # Join each attempt on its own line so a single cell shows every attempt.
    return "\n".join(a_parts), "\n".join(o_parts), "\n".join(d_parts)


def log_session(setup, session_id, attempts, challenge_verdict,
                passive_verdict, passive_live_prob, passive_spoof_prob,
                final_result, final_user_message, device_browser):
    """Append one session row to the Excel workbook, creating it if needed."""

    # Fall back to an empty dict so a missing setup does not raise an error.
    setup = setup or {}

    # Build the three challenge cells from the attempt records.
    actions, outcomes, details = _fmt_attempts(attempts)

    # Assemble the row values in the fixed column order.
    row = [
        session_id or "",                                          # session id
        setup.get("participant_code", ""),                         # participant code
        setup.get("age_group", ""),                                # age band
        setup.get("gender", ""),                                   # gender
        setup.get("ethnicity", ""),                                # ethnicity
        setup.get("visual_aids", ""),                              # visual aids
        datetime.now().strftime("%Y-%m-%d %H:%M:%S"),              # write time
        setup.get("condition", ""),                                # test condition
        actions,                                                   # actions per attempt
        outcomes,                                                  # outcomes per attempt
        details,                                                   # details per attempt
        challenge_verdict or "",                                   # PASSED / FAILED
        passive_verdict or "",                                     # LIVE / SPOOF
        (round(passive_live_prob, 3)                               # live probability
         if passive_live_prob is not None else ""),
        (round(passive_spoof_prob, 3)                              # spoof probability
         if passive_spoof_prob is not None else ""),
        final_result,                                              # combined outcome
        final_user_message,                                        # on-screen message
        device_browser or "",                                      # user-agent string
    ]

    # Ensure the parent folder exists before writing.
    XLSX_PATH.parent.mkdir(parents=True, exist_ok=True)

    try:
        # Open the existing workbook, or create a new one with a header row.
        if XLSX_PATH.exists():
            wb = load_workbook(XLSX_PATH)
            ws = wb.active
        else:
            wb = Workbook()
            ws = wb.active
            ws.title = "sessions"
            ws.append(COLUMNS)
            # Bold the header row.
            for cell in ws[1]:
                cell.font = Font(bold=True)

        # Append this session's row.
        ws.append(row)

        # Wrap text in the three multi-line challenge columns (I, J, K) so the
        # per-attempt lines show on separate lines inside the cell.
        last = ws.max_row
        for col in ("I", "J", "K"):
            ws[f"{col}{last}"].alignment = Alignment(wrap_text=True, vertical="top")

        # Save the workbook.
        wb.save(XLSX_PATH)
        print(f"Session logged: {session_id} | {final_result}")

    except PermissionError:
        # The file is open in Excel and therefore locked.
        print(f"LOG WRITE BLOCKED — the workbook is open in Excel. Close it. "
              f"({XLSX_PATH})")